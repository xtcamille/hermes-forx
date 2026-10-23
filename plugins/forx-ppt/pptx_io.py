"""Native PresentationML export; custom XML stores identity without flattening slides."""
from __future__ import annotations

import hashlib
import io
import json
import zipfile
from copy import deepcopy

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml.xmlchemy import OxmlElement
from pptx.shapes.shapetree import SlideShapes
from pptx.util import Inches, Pt
from pptx.text.text import Font

from .models import effective


META_PATH = "customXml/forx-ppt.xml"
NS = "urn:forx:ppt:project:1"


def _set_color(fill, value):
    if value == "none":
        fill.background()
        return
    fill.solid()
    rgb = fill.fore_color
    if value.startswith("$"):
        rgb.theme_color = {"background": MSO_THEME_COLOR.BACKGROUND_1, "text": MSO_THEME_COLOR.TEXT_1,
                           "muted": MSO_THEME_COLOR.ACCENT_3, "primary": MSO_THEME_COLOR.ACCENT_1,
                           "accent": MSO_THEME_COLOR.ACCENT_2, "surface": MSO_THEME_COLOR.ACCENT_4}[value[1:]]
    else:
        rgb.rgb = RGBColor.from_string(value.lstrip("#"))


def _font(font, style, role="body"):
    if "font_size" in style:
        font.size = Pt(style["font_size"])
    if "bold" in style:
        font.bold = style["bold"]
    if "font_family" in style:
        font.name = style["font_family"]
        rpr = font._rPr
        ea = rpr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
        if ea is None:
            ea = OxmlElement("a:ea")
            rpr.append(ea)
        ea.set("typeface", style["font_family"].replace("-lt", "-ea") if style["font_family"].startswith("+") else style["font_family"])
    if "color" in style:
        _set_color(font.fill, style["color"])


def _text(shape, text, style, role="body"):
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.07)
    frame.margin_top = frame.margin_bottom = Inches(0.04)
    for index, line in enumerate(text.split("\n")):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.line_spacing = 1.2
        p.space_after = Pt(0)
        if "align" in style:
            p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[style["align"]]
        run = p.add_run()
        run.text = line
        _font(run.font, style, role)


def _geometry(shape, box):
    shape.left, shape.top = Inches(box.x), Inches(box.y)
    shape.width, shape.height = Inches(box.width), Inches(box.height)


def _theme(prs, deck):
    part = prs.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(part.blob)
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    mapping = {"dk1": "text", "lt1": "background", "dk2": "text", "lt2": "background",
               "accent1": "primary", "accent2": "accent", "accent3": "muted", "accent4": "surface"}
    for native, key in mapping.items():
        entry = root.find(f".//a:clrScheme/a:{native}", ns)
        for child in list(entry):
            entry.remove(child)
        value = OxmlElement("a:srgbClr")
        value.set("val", deck.theme[key].lstrip("#"))
        entry.append(value)
    for role, key in (("majorFont", "font_title"), ("minorFont", "font_body")):
        for tag in ("latin", "ea", "cs"):
            root.find(f".//a:fontScheme/a:{role}/a:{tag}", ns).set("typeface", deck.theme[key])
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _master_styles(prs, deck):
    master = prs.slide_master
    for role, tag in (("title", "titleStyle"), ("body", "bodyStyle")):
        container = master._element.find(f"{{http://schemas.openxmlformats.org/presentationml/2006/main}}txStyles/{{http://schemas.openxmlformats.org/presentationml/2006/main}}{tag}")
        for child in list(container):
            container.remove(child)
        ppr = OxmlElement("a:lvl1pPr")
        ppr.set("marL", "0")
        ppr.set("indent", "0")
        ppr.append(OxmlElement("a:buNone"))
        rpr = OxmlElement("a:defRPr")
        values = {"font_size": deck.theme["title_size" if role == "title" else "body_size"],
                  "bold": role == "title", "color": "$text"}
        values.update(deck.master.styles.get(role).model_dump(exclude_none=True))
        if "align" in values:
            ppr.set("algn", {"left": "l", "center": "ctr", "right": "r"}[values["align"]])
        _font(Font(rpr), values, role)
        ppr.append(rpr)
        container.append(ppr)
    _set_color(master.background.fill, "$background")


def _layouts(prs, deck):
    seed = prs.slide_layouts[1]
    placeholders = [deepcopy(seed.placeholders[i]._element) for i in (0, 1)]
    while len(prs.slide_layouts) > len(deck.layouts):
        prs.slide_layouts.remove(prs.slide_layouts[-1])
    result = {}
    for native, layout in zip(prs.slide_layouts, deck.layouts):
        for shape in list(native.shapes):
            shape._element.getparent().remove(shape._element)
        native.name = f"{layout.id}_{layout.name}"
        for index, role in enumerate(("title", "body")):
            node = deepcopy(placeholders[index])
            native.shapes._spTree.insert_element_before(node, "p:extLst")
            ph = node.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}ph")
            ph.set("type", role)
            ph.set("idx", str(index))
        for shape, role in zip(native.placeholders, ("title", "body")):
            shape.name = f"{layout.id}_{role}"
            _geometry(shape, layout.boxes[role])
            # Seed layout text must not carry template-specific direct formatting.
            _text(shape, "", {})
            style = layout.styles.get(role)
            if style:
                # Layout paragraph defaults, rather than sample runs, are inherited.
                ppr = shape.text_frame.paragraphs[0]._p.get_or_add_pPr()
                _font(Font(ppr.get_or_add_defRPr()), style.model_dump(exclude_none=True), role)
                if style.align:
                    ppr.set("algn", {"left": "l", "center": "ctr", "right": "r"}[style.align])
        result[layout.id] = native
    return result


def _render_text(shapes, e, box, style, store):
    shape = shapes.add_textbox(Inches(box.x), Inches(box.y), Inches(box.width), Inches(box.height))
    _text(shape, e.text, style)
    _set_color(shape.fill, style["fill"])
    return shape


def _render_shape(shapes, e, box, style, store):
    kind = e.shape or style.pop("default_shape", "rect")
    if kind == "line":
        return shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(box.x), Inches(box.y), Inches(box.x + box.width), Inches(box.y + box.height))
    kinds = {"rect": MSO_AUTO_SHAPE_TYPE.RECTANGLE, "roundRect": MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
             "ellipse": MSO_AUTO_SHAPE_TYPE.OVAL, "line": MSO_AUTO_SHAPE_TYPE.RECTANGLE}
    shape = shapes.add_shape(kinds[kind], Inches(box.x), Inches(box.y), Inches(box.width), Inches(box.height))
    _set_color(shape.fill, style["fill"] if e.style.fill is not None else "$surface")
    return shape


def _render_image(shapes, e, box, style, store):
    shape = shapes.add_picture(io.BytesIO(store.asset(e.asset_id)), Inches(box.x), Inches(box.y), Inches(box.width), Inches(box.height))
    shape.crop_left, shape.crop_top, shape.crop_right, shape.crop_bottom = e.crop
    return shape


def _render_table(shapes, e, box, style, store):
    shape = shapes.add_table(len(e.rows), len(e.rows[0]), Inches(box.x), Inches(box.y), Inches(box.width), Inches(box.height))
    weights = e.column_widths or [1] * len(e.rows[0])
    for column, weight in zip(shape.table.columns, weights):
        column.width = Inches(box.width * weight / sum(weights))
    for r, row in enumerate(e.rows):
        for c, text in enumerate(row):
            cell = shape.table.cell(r, c)
            _set_color(cell.fill, "$primary" if r == 0 else "$surface")
            proxy = type("CellText", (), {"text_frame": cell.text_frame})()
            _text(proxy, text, {**style, "color": "$background" if r == 0 else style["color"], "bold": e.style.bold if e.style.bold is not None else r == 0})
    return shape


def _render_chart(shapes, e, box, style, store):
    data = CategoryChartData()
    data.categories = e.categories
    for series in e.series:
        data.add_series(series.name, series.values)
    kind = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE, "pie": XL_CHART_TYPE.PIE}[e.chart_type]
    shape = shapes.add_chart(kind, Inches(box.x), Inches(box.y), Inches(box.width), Inches(box.height), data)
    chart = shape.chart
    chart.has_title = bool(e.text)
    if e.text:
        _text(type("ChartTitle", (), {"text_frame": chart.chart_title.text_frame})(), e.text, style)
    for i, series in enumerate(chart.series):
        value = ("$primary", "$accent", "$muted")[i % 3]
        _set_color(series.format.fill, value)
        _set_color(series.format.line.fill, value)
        if e.chart_type == "pie":
            for j, point in enumerate(series.points):
                _set_color(point.format.fill, ("$primary", "$accent", "$muted", "$surface")[j % 4])
    chart.has_legend = len(e.series) > 1 or e.chart_type == "pie"
    if chart.has_legend:
        _font(chart.legend.font, {"font_size": 12, "color": "$text"})
    if e.chart_type != "pie":
        for axis in (chart.category_axis, chart.value_axis):
            _font(axis.tick_labels.font, {"font_size": 12, "color": "$text"})
    return shape


RENDERERS = {"text": _render_text, "shape": _render_shape, "image": _render_image,
             "table": _render_table, "chart": _render_chart}


def _render(shapes, deck, slide, e, store):
    if not e.visible:
        return
    style, box = effective(deck, slide, e)
    if e.style.font_family is None:
        style["font_family"] = "+mn-lt"
    style["default_shape"] = deck.theme["shape"]
    if e.kind == "group":
        shape = shapes.add_group_shape()
        for child in e.children:
            _render(shape.shapes, deck, slide, child, store)
    else:
        shape = RENDERERS[e.kind](shapes, e, box, style, store)
    shape.name = f"{e.id}_{e.name}"
    if e.kind in ("text", "shape"):
        _set_color(shape.line.fill, "$primary" if e.shape == "line" and style["line"] == "none" else style["line"])
        shape.line.width = Pt(style["line_width"])


def _digest(blob):
    with zipfile.ZipFile(io.BytesIO(blob)) as archive:
        pairs = [(name, hashlib.sha256(archive.read(name)).hexdigest()) for name in sorted(archive.namelist()) if name != META_PATH]
    return hashlib.sha256(json.dumps(pairs, separators=(",", ":")).encode()).hexdigest()


def export_pptx(deck, store):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(deck.width), Inches(deck.height)
    prs.core_properties.title = deck.name
    prs.core_properties.author = "ForX PPT"
    _theme(prs, deck)
    _master_styles(prs, deck)
    layouts = _layouts(prs, deck)
    master_shapes = SlideShapes(prs.slide_master.shapes._spTree, prs.slide_master)
    for e in deck.master.elements:
        _render(master_shapes, deck, deck.slides[0], e, store)
    for slide in deck.slides:
        native = prs.slides.add_slide(layouts[slide.layout_id])
        native.name = f"{slide.id}_{slide.name}"
        for shape in list(native.shapes):
            shape._element.getparent().remove(shape._element)
        if slide.background:
            _set_color(native.background.fill, slide.background)
        for e in slide.elements:
            if not e.visible:
                continue
            if e.role:
                layout = layouts[slide.layout_id]
                source = next(p for p in layout.placeholders if p.placeholder_format.idx == (0 if e.role == "title" else 1))
                native.shapes.clone_placeholder(source)
                shape = native.shapes[-1]
                shape.name = f"{e.id}_{e.name}"
                if e.box:
                    _geometry(shape, e.box)
                _text(shape, e.text, e.style.model_dump(exclude_none=True), e.role)
                if e.style.fill:
                    _set_color(shape.fill, e.style.fill)
                if e.style.line:
                    _set_color(shape.line.fill, e.style.line)
                    shape.line.width = Pt(e.style.line_width or deck.theme["line_width"])
            else:
                _render(native.shapes, deck, slide, e, store)
        if slide.notes:
            native.notes_slide.notes_text_frame.text = slide.notes
    meta = Part(PackURI("/" + META_PATH), "application/xml", prs.part.package, b"<placeholder/>")
    prs.part.relate_to(meta, RT.CUSTOM_XML)
    output = io.BytesIO()
    prs.save(output)
    root = etree.Element(f"{{{NS}}}project", nsmap={"forx": NS})
    root.set("parts_sha256", _digest(output.getvalue()))
    root.text = json.dumps(store.bundle(deck), ensure_ascii=False, separators=(",", ":"))
    meta._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8")
    output = io.BytesIO()
    prs.save(output)
    return output.getvalue()


def import_pptx(blob):
    try:
        with zipfile.ZipFile(io.BytesIO(blob)) as archive:
            entries = archive.infolist()
            if len(entries) > 20000 or sum(i.file_size for i in entries) > 150 * 1024 * 1024:
                raise ValueError("文件解压后过大")
            if META_PATH not in archive.namelist():
                raise ValueError("首版只支持重新导入本插件导出的 PPTX；也可导入 .forxppt 项目")
            parser = etree.XMLParser(resolve_entities=False, no_network=True)
            root = etree.fromstring(archive.read(META_PATH), parser)
            if root.tag != f"{{{NS}}}project" or root.get("parts_sha256") != _digest(blob):
                raise ValueError("PPTX 已被外部软件修改，首版不能可靠恢复编号；请使用原项目继续编辑")
            return json.loads(root.text)
    except (zipfile.BadZipFile, etree.XMLSyntaxError, json.JSONDecodeError) as exc:
        raise ValueError("PPTX 文件无法读取") from exc
