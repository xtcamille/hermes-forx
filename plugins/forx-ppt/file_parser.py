"""Robust multi-format document parser extracting plain text for presentation synthesis.

Supports .docx, .pptx, .txt, .md, .csv, .json, .html, and text formats.
Zero external dependencies beyond optional python-pptx (standard library zipfile + xml for docx).
"""
from __future__ import annotations

import io
from pathlib import Path
from typing import Any


def extract_text_from_docx(data: bytes) -> str:
    """Extract plain text from .docx document.xml without requiring third-party libraries."""
    import zipfile
    import xml.etree.ElementTree as ET

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if "word/document.xml" not in z.namelist():
                return ""
            xml_content = z.read("word/document.xml")
    except Exception:
        return ""

    root = ET.fromstring(xml_content)
    paragraphs: list[str] = []
    for p in root.iter():
        if p.tag.endswith("}p"):
            p_parts: list[str] = []
            for node in p.iter():
                if node.tag.endswith("}t") and node.text:
                    p_parts.append(node.text)
                elif node.tag.endswith("}tab"):
                    p_parts.append("\t")
                elif node.tag.endswith("}br"):
                    p_parts.append("\n")
            line = "".join(p_parts).strip()
            if line:
                paragraphs.append(line)
    return "\n\n".join(paragraphs)


def extract_text_from_pptx(data: bytes) -> str:
    """Extract text from slides in .pptx presentations."""
    try:
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
    except Exception:
        return ""

    slides_text: list[str] = []
    for idx, slide in enumerate(prs.slides, 1):
        lines: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    t = paragraph.text.strip()
                    if t:
                        lines.append(t)
        if lines:
            slides_text.append(f"【幻灯片 {idx}】\n" + "\n".join(lines))
    return "\n\n".join(slides_text)


def extract_text_from_plain(data: bytes) -> str:
    """Decode raw bytes into string with fallback encodings."""
    for enc in ("utf-8", "utf-8-sig", "gb18030", "gbk", "latin1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def parse_uploaded_file(filename: str, raw: bytes) -> dict[str, Any]:
    """Parse raw file bytes into extracted text and metadata."""
    ext = Path(filename).suffix.lower()

    if ext in (".docx", ".doc"):
        # If it's a real docx zip archive
        text = extract_text_from_docx(raw)
        if not text:
            # Fallback if docx extraction failed or plain text
            text = extract_text_from_plain(raw)
        fmt = "docx"
    elif ext in (".pptx", ".ppt"):
        text = extract_text_from_pptx(raw)
        if not text:
            text = extract_text_from_plain(raw)
        fmt = "pptx"
    elif ext in (".json", ".jsonl"):
        import json
        raw_text = extract_text_from_plain(raw)
        try:
            parsed = json.loads(raw_text)
            text = json.dumps(parsed, ensure_ascii=False, indent=2)
        except Exception:
            text = raw_text
        fmt = "json"
    elif ext in (".csv", ".tsv"):
        text = extract_text_from_plain(raw)
        fmt = "csv"
    else:
        text = extract_text_from_plain(raw)
        fmt = ext.lstrip(".") or "txt"

    clean_text = text.strip()
    return {
        "filename": filename,
        "format": fmt,
        "size": len(raw),
        "char_count": len(clean_text),
        "text": clean_text,
    }
